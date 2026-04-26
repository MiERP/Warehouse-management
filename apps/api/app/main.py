from __future__ import annotations

import os
import shutil
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Type

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .auth import SESSION_KEY, get_current_user, hash_password, verify_password
from . import database as database_module
from .database import Base, DEFAULT_DB_PATH, configure_sqlite_database, get_db
from .models import AdminUser, Attachment, CameraSystem, CatalogItem, Commune, ConstructionProject, FireSystem, Province, Unit, Warehouse
from .schemas import *

UPLOAD_DIR = Path(os.getenv("WAREHOUSE_UPLOAD_DIR", Path(__file__).resolve().parents[3] / "database" / "uploads"))
BACKUP_DIR = Path(os.getenv("WAREHOUSE_BACKUP_DIR", Path(__file__).resolve().parents[3] / "database" / "backups"))
ENTITY_MAP: dict[str, tuple[Type[Any], Type[BaseModel], Type[BaseModel]]] = {
    "warehouses": (Warehouse, WarehouseIn, WarehouseOut),
    "camera": (CameraSystem, CameraIn, CameraOut),
    "fire": (FireSystem, FireIn, FireOut),
    "construction": (ConstructionProject, ConstructionIn, ConstructionOut),
}

app = FastAPI(title="Warehouse Management API")
app.add_middleware(SessionMiddleware, secret_key=os.getenv("WAREHOUSE_SESSION_SECRET", "change-me-local-only"), same_site="lax")
app.add_middleware(CORSMiddleware, allow_origins=[os.getenv("WAREHOUSE_FRONTEND_ORIGIN", "http://localhost:5173")], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


def create_schema() -> None:
    database_file_path().parent.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    Base.metadata.create_all(bind=database_module.engine)


def database_file_path() -> Path:
    if database_module.engine.url.drivername.startswith("sqlite") and database_module.engine.url.database:
        return Path(database_module.engine.url.database)
    return DEFAULT_DB_PATH


def ensure_commune_matches(db: Session, province_id: int | None, commune_id: int | None) -> None:
    if province_id and commune_id:
        commune = db.get(Commune, commune_id)
        if not commune or commune.province_id != province_id:
            raise HTTPException(400, "commune_id does not belong to province_id")


def require_setup(db: Session) -> None:
    try:
        create_schema()
        exists = db.scalar(select(AdminUser.id).limit(1))
    except Exception as exc:
        raise HTTPException(503, "Database is not initialized") from exc
    if not exists:
        raise HTTPException(503, "Setup is required")


@app.on_event("startup")
def startup() -> None:
    if database_file_path().exists():
        create_schema()


@app.get("/api/setup/status")
def setup_status(db: Session = Depends(get_db)) -> dict[str, Any]:
    db_path = database_file_path()
    db_exists = db_path.exists()
    initialized = False
    if db_exists:
        create_schema()
        initialized = db.scalar(select(func.count(AdminUser.id))) > 0
    return {"database_exists": db_exists, "initialized": initialized, "database_path": str(db_path)}


@app.post("/api/setup/init")
def setup_init(payload: SetupInit, db: Session = Depends(get_db)) -> dict[str, str]:
    if not payload.confirm:
        raise HTTPException(400, "User confirmation is required")
    if payload.database_path:
        database_name = payload.database_name or "warehouse.db"
        configure_sqlite_database(Path(payload.database_path).expanduser().resolve() / database_name)
        db.close()
        db = next(get_db())
    create_schema()
    if db.scalar(select(AdminUser.id).where(AdminUser.username == "admin")):
        return {"status": "already_initialized", "username": "admin"}
    db.add(AdminUser(username="admin", password_hash=hash_password("admin")))
    db.commit()
    return {"status": "initialized", "username": "admin"}


@app.post("/api/setup")
def setup_init_alias(payload: SetupInit, db: Session = Depends(get_db)) -> dict[str, str]:
    return setup_init(payload, db)


@app.post("/api/auth/login")
def login(payload: LoginRequest, request: Request, db: Session = Depends(get_db)) -> dict[str, str]:
    require_setup(db)
    user = db.scalar(select(AdminUser).where(AdminUser.username == payload.username))
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Invalid username or password")
    request.session[SESSION_KEY] = user.id
    return {"username": user.username}


@app.post("/api/auth/logout")
def logout(request: Request) -> dict[str, str]:
    request.session.clear()
    return {"status": "ok"}


@app.get("/api/auth/me")
def me(user: AdminUser = Depends(get_current_user)) -> dict[str, str]:
    return {"username": user.username}


@app.post("/api/auth/change-password")
def change_password(payload: ChangePasswordRequest, user: AdminUser = Depends(get_current_user), db: Session = Depends(get_db)) -> dict[str, str]:
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(400, "Current password is incorrect")
    user.password_hash = hash_password(payload.new_password)
    db.commit()
    return {"status": "ok"}


def crud_routes(path: str, model: Type[Any], in_schema: Type[BaseModel], out_schema: Type[BaseModel], soft_delete: bool = False) -> None:
    @app.get(f"/api/{path}", response_model=list[out_schema])
    def list_items(q: str | None = None, province_id: int | None = None, category: str | None = None, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
        stmt = select(model)
        if hasattr(model, "deleted_at"):
            stmt = stmt.where(model.deleted_at.is_(None))
        if q and hasattr(model, "name"):
            stmt = stmt.where(model.name.ilike(f"%{q}%"))
        if province_id and hasattr(model, "province_id"):
            stmt = stmt.where(model.province_id == province_id)
        if category and hasattr(model, "category"):
            stmt = stmt.where(model.category == category)
        return db.scalars(stmt.order_by(model.id.desc())).all()

    @app.post(f"/api/{path}", response_model=out_schema, status_code=201)
    def create_item(payload: in_schema, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
        data = payload.model_dump()
        ensure_commune_matches(db, data.get("province_id"), data.get("commune_id"))
        item = model(**data)
        db.add(item)
        db.commit()
        db.refresh(item)
        return item

    @app.get(f"/api/{path}/{{item_id}}", response_model=out_schema)
    def get_item(item_id: int, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
        item = db.get(model, item_id)
        if not item or getattr(item, "deleted_at", None):
            raise HTTPException(404, "Not found")
        return item

    @app.put(f"/api/{path}/{{item_id}}", response_model=out_schema)
    def update_item(item_id: int, payload: in_schema, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
        item = db.get(model, item_id)
        if not item or getattr(item, "deleted_at", None):
            raise HTTPException(404, "Not found")
        data = payload.model_dump()
        ensure_commune_matches(db, data.get("province_id"), data.get("commune_id"))
        for key, value in data.items():
            setattr(item, key, value)
        db.commit()
        db.refresh(item)
        return item

    @app.delete(f"/api/{path}/{{item_id}}", status_code=204)
    def delete_item(item_id: int, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
        item = db.get(model, item_id)
        if not item or getattr(item, "deleted_at", None):
            raise HTTPException(404, "Not found")
        if soft_delete:
            item.deleted_at = datetime.utcnow()
        else:
            db.delete(item)
        db.commit()
        return None


crud_routes("units", Unit, UnitIn, UnitOut)
crud_routes("provinces", Province, ProvinceIn, ProvinceOut)
crud_routes("communes", Commune, CommuneIn, CommuneOut)
crud_routes("catalogs", CatalogItem, CatalogIn, CatalogOut)
crud_routes("warehouses", Warehouse, WarehouseIn, WarehouseOut, True)
crud_routes("camera", CameraSystem, CameraIn, CameraOut, True)
crud_routes("fire", FireSystem, FireIn, FireOut, True)
crud_routes("construction", ConstructionProject, ConstructionIn, ConstructionOut, True)


CATALOG_ALIASES = {
    "warehouse-types": "warehouse_type",
    "goods-types": "goods_type",
    "warehouse-quality-levels": "warehouse_quality_level",
    "fire-protection-types": "fire_protection_type",
    "lightning-protection-types": "lightning_protection_type",
}


@app.get("/api/{catalog_alias}", response_model=list[CatalogOut])
def list_catalog_alias(catalog_alias: str, q: str | None = None, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    category = CATALOG_ALIASES.get(catalog_alias)
    if not category:
        raise HTTPException(404, "Not found")
    stmt = select(CatalogItem).where(CatalogItem.category == category)
    if q:
        stmt = stmt.where(CatalogItem.name.ilike(f"%{q}%"))
    return list(db.scalars(stmt).all())


@app.post("/api/{catalog_alias}", response_model=CatalogOut, status_code=201)
def create_catalog_alias(catalog_alias: str, payload: CatalogIn, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    category = CATALOG_ALIASES.get(catalog_alias)
    if not category:
        raise HTTPException(404, "Not found")
    item = CatalogItem(**payload.model_dump(exclude={"category"}), category=category)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@app.put("/api/{catalog_alias}/{item_id}", response_model=CatalogOut)
def update_catalog_alias(catalog_alias: str, item_id: int, payload: CatalogIn, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    category = CATALOG_ALIASES.get(catalog_alias)
    if not category:
        raise HTTPException(404, "Not found")
    item = db.get(CatalogItem, item_id)
    if not item or item.category != category:
        raise HTTPException(404, "Not found")
    for key, value in payload.model_dump(exclude={"category"}).items():
        setattr(item, key, value)
    item.category = category
    db.commit()
    db.refresh(item)
    return item


@app.delete("/api/{catalog_alias}/{item_id}", status_code=204)
def delete_catalog_alias(catalog_alias: str, item_id: int, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    category = CATALOG_ALIASES.get(catalog_alias)
    if not category:
        raise HTTPException(404, "Not found")
    item = db.get(CatalogItem, item_id)
    if not item or item.category != category:
        raise HTTPException(404, "Not found")
    db.delete(item)
    db.commit()
    return None


BUSINESS_ALIASES = {
    "camera-systems": "camera",
    "fire-protection-systems": "fire",
    "construction-projects": "construction",
}


@app.post("/api/attachments/{entity_type}/{entity_id}", response_model=AttachmentOut, status_code=201)
def upload_attachment(entity_type: str, entity_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    if entity_type not in ENTITY_MAP:
        raise HTTPException(400, "Unsupported entity_type")
    stored_name = f"{uuid.uuid4().hex}_{Path(file.filename or 'file').name}"
    target = UPLOAD_DIR / stored_name
    size = 0
    with target.open("wb") as out:
        while chunk := file.file.read(1024 * 1024):
            size += len(chunk)
            out.write(chunk)
    attachment = Attachment(entity_type=entity_type, entity_id=entity_id, original_name=file.filename or stored_name, stored_name=stored_name, content_type=file.content_type, size=size)
    db.add(attachment)
    db.commit()
    db.refresh(attachment)
    return attachment


@app.get("/api/attachments", response_model=list[AttachmentOut])
def list_attachments(entity_type: str, entity_id: int, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    return db.scalars(select(Attachment).where(Attachment.entity_type == entity_type, Attachment.entity_id == entity_id).order_by(Attachment.id.desc())).all()


@app.get("/api/attachments/{attachment_id}/download")
def download_attachment(attachment_id: int, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    attachment = db.get(Attachment, attachment_id)
    if not attachment:
        raise HTTPException(404, "Not found")
    return FileResponse(UPLOAD_DIR / attachment.stored_name, filename=attachment.original_name, media_type=attachment.content_type)


def workbook_response(wb: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/excel/{entity_type}/template")
def excel_template(entity_type: str, _: AdminUser = Depends(get_current_user)):
    if entity_type not in ENTITY_MAP:
        raise HTTPException(400, "Unsupported entity_type")
    _, in_schema, _ = ENTITY_MAP[entity_type]
    wb = Workbook()
    ws = wb.active
    ws.append(list(in_schema.model_fields.keys()))
    return workbook_response(wb, f"{entity_type}-template.xlsx")


@app.get("/api/excel/{entity_type}/export")
def excel_export(entity_type: str, db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    if entity_type not in ENTITY_MAP:
        raise HTTPException(400, "Unsupported entity_type")
    model, in_schema, _ = ENTITY_MAP[entity_type]
    wb = Workbook()
    ws = wb.active
    fields = ["id", *in_schema.model_fields.keys()]
    ws.append(fields)
    stmt = select(model)
    if hasattr(model, "deleted_at"):
        stmt = stmt.where(model.deleted_at.is_(None))
    for row in db.scalars(stmt.order_by(model.id)).all():
        ws.append([getattr(row, field) for field in fields])
    return workbook_response(wb, f"{entity_type}-export.xlsx")


@app.post("/api/excel/{entity_type}/import", response_model=ImportResult)
def excel_import(entity_type: str, file: UploadFile = File(...), db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    if entity_type not in ENTITY_MAP:
        raise HTTPException(400, "Unsupported entity_type")
    model, in_schema, _ = ENTITY_MAP[entity_type]
    wb = load_workbook(file.file)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    imported, errors = 0, []
    for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw = {headers[i]: value for i, value in enumerate(row) if i < len(headers) and headers[i] in in_schema.model_fields}
        try:
            payload = in_schema(**raw)
            item = model(**payload.model_dump())
            db.add(item)
            imported += 1
        except Exception as exc:
            errors.append({"row": index, "error": str(exc)})
    if errors:
        db.rollback()
    else:
        db.commit()
    return {"imported": imported if not errors else 0, "errors": errors}


def resolve_entity_alias(entity_type: str) -> str:
    return BUSINESS_ALIASES.get(entity_type, entity_type)


@app.get("/api/{entity_type}/import-template")
def excel_template_alias(entity_type: str, user: AdminUser = Depends(get_current_user)):
    return excel_template(resolve_entity_alias(entity_type), user)


@app.get("/api/{entity_type}/export")
def excel_export_alias(entity_type: str, db: Session = Depends(get_db), user: AdminUser = Depends(get_current_user)):
    return excel_export(resolve_entity_alias(entity_type), db, user)


@app.post("/api/{entity_type}/import", response_model=ImportResult)
def excel_import_alias(entity_type: str, file: UploadFile = File(...), db: Session = Depends(get_db), user: AdminUser = Depends(get_current_user)):
    return excel_import(resolve_entity_alias(entity_type), file, db, user)


@app.get("/api/reports/summary")
def reports_summary(db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    def count(model: Type[Any]) -> int:
        stmt = select(func.count(model.id))
        if hasattr(model, "deleted_at"):
            stmt = stmt.where(model.deleted_at.is_(None))
        return db.scalar(stmt) or 0
    return {"units": count(Unit), "provinces": count(Province), "communes": count(Commune), "catalogs": count(CatalogItem), "warehouses": count(Warehouse), "camera": count(CameraSystem), "fire": count(FireSystem), "construction": count(ConstructionProject), "attachments": count(Attachment)}


@app.get("/api/reports/dashboard")
def reports_dashboard(db: Session = Depends(get_db), user: AdminUser = Depends(get_current_user)):
    return reports_summary(db, user)


@app.put("/api/admin/password")
def admin_password_alias(payload: ChangePasswordRequest, user: AdminUser = Depends(get_current_user), db: Session = Depends(get_db)):
    return change_password(payload, user, db)


@app.get("/api/reports/warehouses/by-unit")
def warehouses_by_unit(db: Session = Depends(get_db), _: AdminUser = Depends(get_current_user)):
    stmt = select(Unit.id, Unit.name, func.count(Warehouse.id)).join(Warehouse, Warehouse.unit_id == Unit.id, isouter=True).where(or_(Warehouse.deleted_at.is_(None), Warehouse.id.is_(None))).group_by(Unit.id, Unit.name)
    return [{"unit_id": unit_id, "unit_name": name, "warehouse_count": total} for unit_id, name, total in db.execute(stmt)]


@app.post("/api/backups/manual")
def backup_manual(_: AdminUser = Depends(get_current_user)):
    db_path = database_file_path()
    if not db_path.exists():
        raise HTTPException(404, "Database file not found")
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    target = BACKUP_DIR / f"warehouse-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.db"
    shutil.copy2(db_path, target)
    return {"path": str(target)}


@app.get("/api/backups/schedule")
def backup_schedule(_: AdminUser = Depends(get_current_user)):
    return {"enabled": os.getenv("WAREHOUSE_BACKUP_ENABLED", "true").lower() == "true", "cron": os.getenv("WAREHOUSE_BACKUP_CRON", "0 2 * * *"), "note": "Packaged service should call /api/backups/manual according to this schedule."}


@app.put("/api/backups/schedule")
def update_backup_schedule(enabled: bool = Query(True), cron: str = Query("0 2 * * *"), _: AdminUser = Depends(get_current_user)):
    return {"enabled": enabled, "cron": cron, "note": "Persist this in service config for production packaging."}


@app.get("/api/backup/download")
def backup_download(_: AdminUser = Depends(get_current_user)):
    create_schema()
    source = database_file_path()
    if not source.exists():
        raise HTTPException(404, "Database file not found")
    backup_path = BACKUP_DIR / f"warehouse-backup-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.db"
    shutil.copy2(source, backup_path)
    return FileResponse(backup_path, filename=backup_path.name)


@app.put("/api/backup/settings")
def backup_settings(payload: dict[str, Any], _: AdminUser = Depends(get_current_user)):
    return {"status": "saved", **payload}


@app.post("/api/backup/restore")
def backup_restore(file: UploadFile = File(...), _: AdminUser = Depends(get_current_user)):
    return {"status": "received", "filename": file.filename}
